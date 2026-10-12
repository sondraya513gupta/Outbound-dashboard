from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import LEADS_CSV, LEADS_JSON, LEADS_XLSX
from models import Campaign, Lead, Store
from services.search import plausible_email

COLUMNS = [
    ("organization_name", "Organization Name"),
    ("website", "Website"),
    ("country", "Country"),
    ("industry", "Industry"),
    ("contact_email", "Contact Email"),
    ("description", "Description"),
    ("source_url", "Source URL"),
    ("email_subject", "Email Subject"),
    ("email_body", "Email Body"),
    ("status", "Status"),
]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def load_store() -> Store:
    if not LEADS_JSON.exists():
        return Store()
    return Store.model_validate_json(LEADS_JSON.read_text(encoding="utf-8"))


def save_store(store: Store) -> None:
    LEADS_JSON.write_text(
        store.model_dump_json(indent=2),
        encoding="utf-8",
    )
    _write_csv(store.leads)
    _write_xlsx(store.leads)


def add_campaign(payload, leads_raw: list[dict]) -> tuple[Campaign, list[Lead]]:
    store = load_store()
    campaign = Campaign(
        id=str(uuid4())[:8],
        product_description=payload.product_description,
        country=payload.country,
        industry=payload.industry,
        requested=payload.number_of_leads,
        created_at=_now(),
    )
    leads: list[Lead] = []
    for item in leads_raw:
        lead = Lead(
            id=str(uuid4())[:8],
            campaign_id=campaign.id,
            organization_name=item["organization_name"],
            website=item["website"],
            country=item.get("country") or payload.country,
            industry=item.get("industry") or payload.industry,
            contact_email=plausible_email(item.get("contact_email")),
            description=item.get("description") or "",
            source_url=item.get("source_url") or item["website"],
            email_subject=item.get("email_subject") or "",
            email_body=item.get("email_body") or "",
            status="Not Contacted",
        )
        leads.append(lead)
        campaign.lead_ids.append(lead.id)
    store.campaigns.insert(0, campaign)
    store.leads = leads + store.leads
    save_store(store)
    return campaign, leads


def list_leads(campaign_id: str | None = None) -> list[Lead]:
    store = load_store()
    leads = store.leads
    if campaign_id:
        leads = [lead for lead in leads if lead.campaign_id == campaign_id]
    return leads


def get_lead(lead_id: str) -> Lead | None:
    for lead in load_store().leads:
        if lead.id == lead_id:
            return lead
    return None


def mark_contacted(lead_id: str) -> Lead | None:
    store = load_store()
    updated = None
    for lead in store.leads:
        if lead.id == lead_id:
            lead.status = "Contacted"
            updated = lead
            break
    if updated:
        save_store(store)
    return updated


def mailto_link(lead: Lead) -> str:
    to = lead.contact_email or ""
    subject = quote(lead.email_subject or "")
    body = quote(lead.email_body or "")
    return f"mailto:{to}?subject={subject}&body={body}"


def _write_csv(leads: list[Lead]) -> None:
    with LEADS_CSV.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in COLUMNS])
        for lead in leads:
            writer.writerow([_cell(lead, key) for key, _ in COLUMNS])


def _write_xlsx(leads: list[Lead]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Leads"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A2E")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, (_, label) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, col, label)
        cell.font = header_font
        cell.fill = header_fill

    for row, lead in enumerate(leads, start=2):
        for col, (key, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row, col, _cell(lead, key))
            cell.alignment = wrap

    widths = [28, 32, 16, 22, 28, 48, 36, 36, 60, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(leads) + 1, 1)}"
    book.save(LEADS_XLSX)


def _cell(lead: Lead, key: str) -> str:
    value = getattr(lead, key)
    return "" if value is None else str(value)


def export_paths() -> dict[str, Path]:
    store = load_store()
    if store.leads:
        save_store(store)
    return {"csv": LEADS_CSV, "xlsx": LEADS_XLSX}
